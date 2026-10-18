"""表格清洗工具的离线验收测试。所有样本与输出均写入临时目录。"""

from __future__ import annotations

import copy
import csv
from datetime import date, datetime
from decimal import Decimal
import hashlib
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook


MODULE_PATH = Path(__file__).with_name("9077_表格清洗.py")
SPEC = importlib.util.spec_from_file_location("cleaner_9077", MODULE_PATH)
cleaner = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = cleaner
SPEC.loader.exec_module(cleaner)

HEADERS = ["order_id", "email", "amount", "order_date", "status", "region"]
BASE_CONFIG = {
    "columns": {
        "order_id": {"type": "string", "required": True, "case": "upper"},
        "email": {"type": "email", "required": True, "case": "lower"},
        "amount": {"type": "decimal", "required": True, "min": "0", "decimal_places": 2},
        "order_date": {"type": "date", "required": True, "formats": ["%Y-%m-%d", "%m/%d/%Y"]},
        "status": {"type": "enum", "required": True, "case": "lower", "choices": ["paid", "pending", "cancelled"]},
        "region": {"type": "string", "case": "upper", "default": "UNKNOWN"},
    },
    "deduplicate_by": ["order_id"],
    "filters": [{"column": "status", "operator": "in", "value": ["paid", "pending"]}],
    "sort": [{"column": "order_date", "descending": False}],
}


class CleanerAcceptanceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="cleaner9077_")
        self.root = Path(self.temp.name)
        self.addCleanup(self.temp.cleanup)

    def write_config(self, config=None, name="规则.json"):
        path = self.root / name
        path.write_text(json.dumps(BASE_CONFIG if config is None else config, ensure_ascii=False), encoding="utf-8")
        return path

    def write_csv(self, rows, headers=None, name="输入.csv"):
        path = self.root / name
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(HEADERS if headers is None else headers)
            writer.writerows(rows)
        return path

    def write_xlsx(self, rows, headers=None, name="输入.xlsx"):
        path = self.root / name
        book = Workbook()
        sheet = book.active
        sheet.append(HEADERS if headers is None else headers)
        for row in rows:
            sheet.append(row)
        book.save(path)
        book.close()
        return path

    @staticmethod
    def source_hash(path):
        return hashlib.sha256(Path(path).read_bytes()).hexdigest()

    @staticmethod
    def read_csv(path):
        with Path(path).open(encoding="utf-8-sig", newline="") as stream:
            return list(csv.DictReader(stream))

    @staticmethod
    def read_xlsx(path):
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            rows = list(book["Clean data"].values)
            records = []
            for row in rows[1:]:
                record = dict(zip(rows[0], row))
                for key, value in record.items():
                    if isinstance(value, (datetime, date)):
                        record[key] = value.strftime("%Y-%m-%d")
                    elif key == "amount" and value is not None:
                        record[key] = format(Decimal(str(value)), ".2f")
                records.append(record)
            return records
        finally:
            book.close()

    def execute(self, source, config=None, output_name="输出"):
        output = self.root / output_name
        result = cleaner.run(source, self.write_config(config), output)
        paths = {
            "csv": output / "9077_清洗结果.csv",
            "xlsx": output / "9077_清洗结果.xlsx",
            "json": output / "9077_审计报告.json",
            "html": output / "9077_审计报告.html",
        }
        for path in paths.values():
            self.assertTrue(path.is_file(), f"缺少输出文件：{path.name}")
        report = json.loads(paths["json"].read_text(encoding="utf-8"))
        return result, paths, report

    def assert_rejected_without_outputs(self, source, config=None):
        output = self.root / "拒绝输出"
        digest = self.source_hash(source)
        with self.assertRaises(cleaner.CleanerError):
            cleaner.run(source, self.write_config(config), output)
        self.assertEqual(digest, self.source_hash(source), "源文件被修改")
        self.assertFalse(output.exists() and any(output.iterdir()), "失败运行留下了输出文件")

    @staticmethod
    def mixed_rows():
        return [
            [" a-002 ", " TWO@EXAMPLE.COM ", " 20.5 ", "10/02/2026", " PENDING ", " east "],
            [" a-001 ", " ONE@EXAMPLE.COM ", "10", "2026-10-01", " PAID ", " "],
            [" A-001 ", "duplicate@example.com", "99", "2026-10-03", "paid", "NORTH"],
            ["cancel-1", "cancel@example.com", "5", "2026-10-01", "cancelled", "west"],
            ["bad-num", "number@example.com", "not money", "2026-10-01", "paid", "east"],
            ["bad-nan", "nan@example.com", "NaN", "2026-10-01", "paid", "east"],
            ["bad-date", "date@example.com", "5", "2026-02-30", "paid", "east"],
            ["no-email", " ", "5", "2026-10-01", "paid", "east"],
        ]

    def test_normalization_defaults_sorting_and_source_preservation(self):
        source = self.write_csv(self.mixed_rows())
        digest = self.source_hash(source)
        _, paths, _ = self.execute(source)
        self.assertEqual(self.read_csv(paths["csv"]), [
            {"order_id": "A-001", "email": "one@example.com", "amount": "10.00", "order_date": "2026-10-01", "status": "paid", "region": "UNKNOWN"},
            {"order_id": "A-002", "email": "two@example.com", "amount": "20.50", "order_date": "2026-10-02", "status": "pending", "region": "EAST"},
        ])
        self.assertEqual(self.source_hash(source), digest)

    def test_csv_and_xlsx_inputs_produce_identical_records(self):
        csv_source = self.write_csv(self.mixed_rows())
        xlsx_source = self.write_xlsx(self.mixed_rows())
        xlsx_hash = self.source_hash(xlsx_source)
        _, csv_paths, _ = self.execute(csv_source, output_name="CSV输出")
        _, xlsx_paths, _ = self.execute(xlsx_source, output_name="XLSX输出")
        expected = self.read_csv(csv_paths["csv"])
        self.assertEqual(self.read_csv(xlsx_paths["csv"]), expected)
        self.assertEqual(self.read_xlsx(csv_paths["xlsx"]), expected)
        self.assertEqual(self.read_xlsx(xlsx_paths["xlsx"]), expected)
        self.assertEqual(xlsx_hash, self.source_hash(xlsx_source))

    def test_actual_workbook_formulas_are_never_evaluated_or_exported(self):
        source = self.write_xlsx([["a-1", "a@example.com", "=1+1", "2026-10-01", "paid", "east"]])
        digest = self.source_hash(source)
        _, paths, report = self.execute(source)
        self.assertEqual(report["counts"]["valid"], 0)
        self.assertEqual(report["counts"]["invalid"], 1)
        self.assertEqual(self.read_csv(paths["csv"]), [])
        self.assertEqual(report["rows"][0]["original"]["amount"], "=1+1")
        self.assertTrue(any("formula" in reason.lower() for reason in report["rows"][0]["reasons"]))
        book = load_workbook(paths["xlsx"], data_only=False)
        try:
            self.assertFalse(any(cell.data_type == "f" for sheet in book for row in sheet for cell in row))
        finally:
            book.close()
        self.assertEqual(digest, self.source_hash(source))

    def test_duplicate_csv_headers_are_rejected(self):
        source = self.write_csv([], headers=["order_id", "email", "amount", "order_date", "status", "order_id"])
        self.assert_rejected_without_outputs(source)

    def test_missing_required_header_is_rejected(self):
        source = self.write_csv([], headers=["order_id", "email", "order_date", "status", "region"])
        self.assert_rejected_without_outputs(source)

    def test_blank_csv_header_is_rejected(self):
        source = self.write_csv([], headers=HEADERS + [""])
        self.assert_rejected_without_outputs(source)

    def test_invalid_config_is_rejected_without_outputs(self):
        source = self.write_csv([["a-1", "a@example.com", "5", "2026-10-01", "paid", "east"]])
        config = copy.deepcopy(BASE_CONFIG)
        config["columns"]["amount"]["type"] = "imaginary_number"
        self.assert_rejected_without_outputs(source, config)

    def test_existing_output_is_not_overwritten_and_no_partial_files(self):
        source = self.write_csv([["a-1", "a@example.com", "5", "2026-10-01", "paid", "east"]])
        output = self.root / "已有输出"
        output.mkdir()
        existing = output / "9077_审计报告.html"
        existing.write_text("保留原报告", encoding="utf-8")
        digest = self.source_hash(source)
        with self.assertRaises(cleaner.OutputError):
            cleaner.run(source, self.write_config(), output)
        self.assertEqual(existing.read_text(encoding="utf-8"), "保留原报告")
        self.assertEqual(sorted(path.name for path in output.iterdir()), [existing.name])
        self.assertEqual(self.source_hash(source), digest)

    def test_audit_accounts_for_every_nonblank_row(self):
        source = self.write_csv(self.mixed_rows() + [["", "", "", "", "", ""]])
        result, paths, report = self.execute(source)
        self.assertEqual(report["counts"], {
            "processed": 8, "valid": 2, "invalid": 4,
            "duplicate": 1, "filtered": 1, "blank_rows_skipped": 1,
        })
        self.assertEqual(result, report)
        self.assertEqual(report["counts"]["processed"], sum(report["counts"][key] for key in ("valid", "invalid", "duplicate", "filtered")))
        self.assertEqual([row["status"] for row in report["rows"]], ["valid", "valid", "duplicate", "filtered", "invalid", "invalid", "invalid", "invalid"])
        self.assertEqual([row["source_row"] for row in report["rows"]], list(range(2, 10)))
        self.assertEqual(report["rows"][0]["original"]["order_id"], " a-002 ")
        self.assertEqual(report["rows"][0]["cleaned"]["order_id"], "A-002")
        for row in report["rows"]:
            if row["status"] != "valid":
                self.assertTrue(row["reasons"], "被排除的数据行缺少原因")
        for kind, path in paths.items():
            self.assertEqual(report["outputs"][kind], path.name)
        self.assertEqual(report["source"], source.name)
        self.assertNotIn(str(self.root), paths["json"].read_text(encoding="utf-8"))

    def test_formula_like_text_is_escaped_in_csv_and_workbook(self):
        values = ["=2+2", "+SUM(1,2)", "-1+2", "@SUM(A1)"]
        rows = [[value, "safe@example.com", "1", "2026-10-01", "paid", "east"] for value in values]
        source = self.write_csv(rows)
        _, paths, report = self.execute(source)
        self.assertEqual(report["counts"]["valid"], 4)
        csv_records = self.read_csv(paths["csv"])
        self.assertEqual([row["order_id"] for row in csv_records], ["'" + value for value in values])
        self.assertEqual([row["cleaned"]["order_id"] for row in report["rows"]], values)
        book = load_workbook(paths["xlsx"], data_only=False)
        try:
            cells = list(book["Clean data"].iter_rows(min_row=2, max_col=1))
            self.assertEqual([row[0].value for row in cells], ["'" + value for value in values])
            self.assertFalse(any(cell.data_type == "f" for sheet in book for row in sheet for cell in row))
        finally:
            book.close()

    def test_formula_like_header_is_safe_in_both_exports(self):
        config = {"columns": {"=1+1": {"type": "string"}}}
        source = self.write_csv([["plain text"]], headers=["=1+1"])
        _, paths, report = self.execute(source, config)
        self.assertEqual(report["counts"]["valid"], 1)
        with paths["csv"].open(encoding="utf-8-sig", newline="") as stream:
            self.assertEqual(next(csv.reader(stream)), ["'=1+1"])
        book = load_workbook(paths["xlsx"], data_only=False)
        try:
            cell = book["Clean data"]["A1"]
            self.assertEqual(cell.value, "'=1+1")
            self.assertNotEqual(cell.data_type, "f")
        finally:
            book.close()

    def test_report_escapes_untrusted_html(self):
        payload = '<img src=x onerror="alert(9077)"><script>alert(9077)</script>'
        source = self.write_csv([[payload, "invalid-email", "1", "2026-10-01", "paid", "east"]])
        _, paths, report = self.execute(source)
        document = paths["html"].read_text(encoding="utf-8")
        self.assertEqual(report["rows"][0]["original"]["order_id"], payload)
        self.assertNotIn(payload, document)
        self.assertNotIn("<img src=x", document)
        self.assertNotIn("<script>alert(9077)", document)
        self.assertIn("&lt;img src=x", document)
        self.assertIn("&lt;script&gt;alert(9077)&lt;/script&gt;", document)

    def test_nonfinite_negative_and_unknown_values_are_invalid(self):
        cases = [
            ["nan", "ok@example.com", "NaN", "2026-10-01", "paid", "east"],
            ["infinity", "ok@example.com", "Infinity", "2026-10-01", "paid", "east"],
            ["negative", "ok@example.com", "-0.01", "2026-10-01", "paid", "east"],
            ["status", "ok@example.com", "1", "2026-10-01", "unknown", "east"],
            ["email", "missing-at-sign", "1", "2026-10-01", "paid", "east"],
        ]
        _, paths, report = self.execute(self.write_csv(cases))
        self.assertEqual(report["counts"]["processed"], 5)
        self.assertEqual(report["counts"]["invalid"], 5)
        self.assertEqual(report["counts"]["valid"], 0)
        self.assertEqual(self.read_csv(paths["csv"]), [])

    def test_decimal_rounding_and_large_values_keep_exact_text(self):
        rows = [
            ["round", "ok@example.com", "1.235", "2026-10-01", "paid", "华东"],
            ["large", "ok@example.com", "123456789012345678.90", "2026-10-01", "paid", "华南"],
        ]
        _, paths, report = self.execute(self.write_csv(rows))
        records = self.read_csv(paths["csv"])
        self.assertEqual([row["amount"] for row in records], ["1.24", "123456789012345678.90"])
        self.assertEqual([row["region"] for row in records], ["华东", "华南"])
        self.assertEqual(report["counts"]["valid"], 2)
        book = load_workbook(paths["xlsx"], data_only=False)
        try:
            sheet = book["Clean data"]
            self.assertEqual(sheet["C3"].value, "123456789012345678.90")
            self.assertEqual(sheet["C3"].data_type, "s")
        finally:
            book.close()

    def test_bad_json_config_leaves_no_outputs(self):
        source = self.write_csv([])
        config = self.root / "损坏规则.json"
        config.write_text("{not valid json", encoding="utf-8")
        output = self.root / "损坏规则输出"
        with self.assertRaises(cleaner.ConfigError):
            cleaner.run(source, config, output)
        self.assertFalse(output.exists())


if __name__ == "__main__":
    unittest.main(verbosity=2)
