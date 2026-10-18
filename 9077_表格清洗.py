"""Local, auditable CSV/XLSX cleaner. Python 3.10+; pip install openpyxl.

Usage: python 9077_表格清洗.py INPUT CONFIG OUTPUT_DIRECTORY
run(input_path, config_path, output_dir) returns the completed audit report.
"""
from __future__ import annotations

import argparse
import csv
import html
import json
import math
import os
import re
import shutil
import sys
import tempfile
from collections import Counter
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, localcontext
from pathlib import Path

PREFIX = "9077_"
NAMES = {"csv": PREFIX + "清洗结果.csv", "xlsx": PREFIX + "清洗结果.xlsx",
         "json": PREFIX + "审计报告.json", "html": PREFIX + "审计报告.html"}
STATUSES = ("valid", "invalid", "duplicate", "filtered")
EMAIL = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")
ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


class CleanerError(Exception):
    """An actionable problem, with a stable CLI exit code."""
    exit_code = 5


class ConfigError(CleanerError):
    exit_code = 2


class InputError(CleanerError):
    exit_code = 3


class OutputError(CleanerError):
    exit_code = 4


def serial(value):
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def safe_cell(value):
    """Escape spreadsheet expression prefixes in text; retain numeric types."""
    value = serial(value)
    if isinstance(value, str) and value.lstrip(" \t\r\n\v\f\ufeff").startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def json_value(value):
    if isinstance(value, dict):
        return {key: json_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_value(item) for item in value]
    return serial(value)


def load_config(path):
    try:
        config = json.loads(Path(path).read_text(encoding="utf-8-sig"))
        if not isinstance(config, dict) or not isinstance(config.get("columns"), dict) or not config["columns"]:
            raise ValueError("columns must be a non-empty object")
        allowed = {"columns", "sheet", "deduplicate_by", "filters", "sort"}
        if set(config) - allowed:
            raise ValueError("unknown configuration key: " + ", ".join(sorted(set(config) - allowed)))
        for name, rule in config["columns"].items():
            if not isinstance(name, str) or not name.strip() or name != name.strip() or ILLEGAL_XML.search(name) or not isinstance(rule, dict):
                raise ValueError("column names must be non-empty, trimmed strings with object rules")
            if rule.get("type", "string") not in {"string", "decimal", "date", "enum", "email"}:
                raise ValueError(f"{name}: unsupported type")
            keys = {"type", "required", "case", "default", "formats", "choices", "min", "max", "decimal_places"}
            if set(rule) - keys:
                raise ValueError(f"{name}: unknown rule key")
            if rule.get("case", "none") not in {"none", "lower", "upper", "title"}:
                raise ValueError(f"{name}: case must be none/lower/upper/title")
            if "required" in rule and not isinstance(rule["required"], bool):
                raise ValueError(f"{name}: required must be boolean")
            places = rule.get("decimal_places", 2)
            if type(places) is not int or not 0 <= places <= 12:
                raise ValueError(f"{name}: decimal_places must be an integer from 0 to 12")
            if rule.get("type") == "date":
                formats = rule.get("formats")
                if not isinstance(formats, list) or not formats or not all(isinstance(x, str) and x for x in formats):
                    raise ValueError(f"{name}: date requires an explicit formats list")
            if rule.get("type") == "enum":
                choices = rule.get("choices")
                if not isinstance(choices, list) or not choices or not all(isinstance(x, str) for x in choices):
                    raise ValueError(f"{name}: enum requires a non-empty string choices list")
            for bound in ("min", "max"):
                if bound in rule and not Decimal(str(rule[bound])).is_finite():
                    raise ValueError(f"{name}: {bound} must be finite")
            if "min" in rule and "max" in rule and Decimal(str(rule["min"])) > Decimal(str(rule["max"])):
                raise ValueError(f"{name}: min exceeds max")
            if "default" in rule:
                clean_value(rule["default"], {k: v for k, v in rule.items() if k != "default"})
        names = set(config["columns"])
        dedup = config.get("deduplicate_by", [])
        if not isinstance(dedup, list) or len(dedup) != len(set(dedup)) or not set(dedup) <= names:
            raise ValueError("deduplicate_by must list unique configured column names")
        for key in ("filters", "sort"):
            if not isinstance(config.get(key, []), list):
                raise ValueError(f"{key} must be a list")
        for item in config.get("filters", []):
            if not isinstance(item, dict) or item.get("column") not in names or item.get("operator") not in {"eq", "ne", "in", "not_in", "gt", "gte", "lt", "lte", "contains"} or "value" not in item:
                raise ValueError("each filter needs column, supported operator, and value")
            if item["operator"] in {"in", "not_in"} and not isinstance(item["value"], list):
                raise ValueError("in/not_in filter value must be a list")
            values = item["value"] if item["operator"] in {"in", "not_in"} else [item["value"]]
            item["_values"] = [clean_value(v, {**config["columns"][item["column"]], "required": False}) for v in values]
        for item in config.get("sort", []):
            if not isinstance(item, dict) or item.get("column") not in names or type(item.get("descending", False)) is not bool:
                raise ValueError("each sort item needs a configured column and optional boolean descending")
        if "sheet" in config and not isinstance(config["sheet"], str):
            raise ValueError("sheet must be a worksheet name")
        return config
    except (OSError, ValueError, TypeError, KeyError, InvalidOperation) as exc:
        raise ConfigError(f"Invalid configuration: {exc}") from exc


def clean_value(value, rule):
    if isinstance(value, str):
        if ILLEGAL_XML.search(value):
            raise ValueError("contains unsupported control characters")
        value = value.strip()
    if value is None or value == "":
        value = rule.get("default")
        if value is None or value == "":
            if rule.get("required", False):
                raise ValueError("required value is missing")
            return None
    kind = rule.get("type", "string")
    if kind == "decimal":
        if isinstance(value, bool):
            raise ValueError("expected a finite decimal number")
        try:
            number = Decimal(str(value))
            if not number.is_finite() or len(number.as_tuple().digits) > 100 or abs(number.adjusted()) > 100:
                raise ValueError("expected a finite decimal with at most 100 digits")
            if "min" in rule and number < Decimal(str(rule["min"])):
                raise ValueError(f"number is below {rule['min']}")
            if "max" in rule and number > Decimal(str(rule["max"])):
                raise ValueError(f"number is above {rule['max']}")
            with localcontext() as ctx:
                ctx.prec = 120
                return number.quantize(Decimal(1).scaleb(-rule.get("decimal_places", 2)), rounding=ROUND_HALF_UP)
        except InvalidOperation as exc:
            raise ValueError("expected a finite decimal number") from exc
    if kind == "date":
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        for pattern in rule["formats"]:
            try:
                return datetime.strptime(str(value), pattern).date()
            except ValueError:
                pass
        raise ValueError("date does not match " + ", ".join(rule["formats"]))
    text = str(value).strip()
    transform = rule.get("case", "none")
    if transform != "none":
        text = getattr(text, transform)()
    if kind == "email" and not EMAIL.fullmatch(text):
        raise ValueError("invalid email address")
    if kind == "enum":
        choices = [getattr(x.strip(), transform)() if transform != "none" else x.strip() for x in rule["choices"]]
        if text not in choices:
            raise ValueError("value must be one of: " + ", ".join(choices))
    return text


def read_source(path, config):
    path = Path(path)
    if not path.is_file():
        raise InputError(f"Input file does not exist: {path}")
    rows, blank = [], 0
    workbook = None
    try:
        if path.suffix.lower() == ".csv":
            with path.open("r", encoding="utf-8-sig", newline="") as stream:
                sample = stream.read(8192)
                stream.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
                records = list(csv.reader(stream, dialect, strict=True))
            if not records:
                raise InputError("Input has no header row")
            raw_headers, data = records[0], ((n, row, []) for n, row in enumerate(records[1:], 2))
        elif path.suffix.lower() == ".xlsx":
            import openpyxl
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
            sheet = workbook[config["sheet"]] if config.get("sheet") else workbook.worksheets[0]
            iterator = sheet.iter_rows()
            first = next(iterator, None)
            if first is None:
                raise InputError("Input has no header row")
            if any(c.data_type in {"f", "e"} for c in first):
                raise InputError("Header row contains a formula or Excel error")
            raw_headers = [c.value for c in first]
            data = ((n, [c.value for c in row], [f"{c.coordinate}: formula cells are not evaluated" if c.data_type == "f" else f"{c.coordinate}: Excel error {c.value}" for c in row if c.data_type in {"f", "e"}]) for n, row in enumerate(iterator, 2))
        else:
            raise InputError("Input must be a UTF-8 CSV or an XLSX file")
        headers = [str(v).strip() if v is not None else "" for v in raw_headers]
        if not headers or any(not h or ILLEGAL_XML.search(h) for h in headers):
            raise InputError("Header names cannot be empty")
        if len({h.casefold() for h in headers}) != len(headers):
            raise InputError("Duplicate header names are not allowed (case-insensitive)")
        missing = set(config["columns"]) - set(headers)
        if missing:
            raise InputError("Missing configured columns: " + ", ".join(sorted(missing)))
        for number, values, reasons in data:
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
                blank += 1
                continue
            original = dict(zip(headers, values + [None] * max(0, len(headers) - len(values))))
            if len(values) > len(headers):
                reasons.append("row has more fields than the header")
                original["__extra_fields__"] = values[len(headers):]
            cleaned = {}
            for name, rule in config["columns"].items():
                try:
                    cleaned[name] = clean_value(original.get(name), rule)
                except (ValueError, TypeError) as exc:
                    reasons.append(f"{name}: {exc}")
            rows.append({"source_row": number, "status": "invalid" if reasons else "valid", "reasons": reasons,
                         "original": original, "cleaned": cleaned})
        return headers, rows, blank
    except InputError:
        raise
    except ImportError as exc:
        raise CleanerError("Missing dependency: install openpyxl with python -m pip install openpyxl") from exc
    except Exception as exc:
        raise InputError(f"Cannot read input: {exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()


def filter_matches(value, item):
    op, targets = item["operator"], item["_values"]
    if op == "in":
        return value in targets
    if op == "not_in":
        return value not in targets
    other = targets[0]
    if op == "eq":
        return value == other
    if op == "ne":
        return value != other
    if value is None or other is None:
        return False
    if op == "contains":
        return str(other) in str(value)
    return {"gt": lambda: value > other, "gte": lambda: value >= other,
            "lt": lambda: value < other, "lte": lambda: value <= other}[op]()


def classify(rows, config):
    seen, keys = {}, config.get("deduplicate_by", [])
    for row in rows:
        if row["status"] != "valid":
            continue
        if keys:
            key = tuple(row["cleaned"][name] for name in keys)
            if key in seen:
                row.update(status="duplicate", reasons=[f"duplicate of source row {seen[key]} on " + ", ".join(keys)])
                continue
            seen[key] = row["source_row"]
        for item in config.get("filters", []):
            if not filter_matches(row["cleaned"][item["column"]], item):
                row.update(status="filtered", reasons=[f"excluded by filter: {item['column']} {item['operator']} {item['value']}"])
                break
    valid = [row for row in rows if row["status"] == "valid"]
    for item in reversed(config.get("sort", [])):
        name = item["column"]
        present = [row for row in valid if row["cleaned"][name] is not None]
        missing = [row for row in valid if row["cleaned"][name] is None]
        valid = sorted(present, key=lambda row: row["cleaned"][name], reverse=item.get("descending", False)) + missing
    return valid


def write_workbook(path, columns, valid, report):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    book = openpyxl.Workbook()
    summary = book.active
    summary.title = "Summary"
    summary.append(["LOCAL DATA CLEANER", "AUDIT SUMMARY"])
    summary.append(["Source", safe_cell(report["source"])])
    summary.append(["Generated UTC", report["generated_at"]])
    for name, count in report["counts"].items():
        summary.append([name.replace("_", " ").title(), count])
    summary.append(["Reconciliation", "processed = valid + invalid + duplicate + filtered"])
    summary.append(["Formula policy", "Source formulas are never evaluated; their rows are invalid."])
    summary.append(["Deduplication", "First valid normalized key wins before filtering."])
    data = book.create_sheet("Clean data")
    data.append([safe_cell(name) for name in columns])
    for row in valid:
        values = []
        for name in columns:
            value = row["cleaned"][name]
            if isinstance(value, Decimal):
                # Excel stores 15 significant digits; longer numbers remain exact text.
                value = float(value) if len(value.as_tuple().digits) <= 15 else str(value)
            values.append(safe_cell(value) if not isinstance(value, date) else value)
        data.append(values)
    review = book.create_sheet("Review")
    review.append(["Source row", "Status", "Reasons", "Original values (JSON)"])
    for row in report["rows"]:
        review.append([row["source_row"], row["status"], safe_cell("; ".join(row["reasons"])),
                       safe_cell(json.dumps(row["original"], ensure_ascii=False, default=serial))])
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 27
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="143347")
            cell.font = Font(color="FFFFFF", bold=True)
        for index, col in enumerate(sheet.columns, 1):
            width = min(68, max(16, max(len(str(c.value or "")) for c in col) + 2))
            sheet.column_dimensions[get_column_letter(index)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F0F5F8")
                if isinstance(cell.value, (date, datetime)):
                    cell.number_format = "yyyy-mm-dd"
    for index, (name, rule) in enumerate(columns.items(), 1):
        if rule.get("type") == "decimal":
            for cells in data.iter_cols(min_col=index, max_col=index, min_row=2):
                for cell in cells:
                    cell.number_format = "#,##0" + ("." + "0" * rule.get("decimal_places", 2) if rule.get("decimal_places", 2) else "")
    book.save(path)
    book.close()


def write_html(path, report):
    esc = lambda v: html.escape(str(v), quote=True)
    cards = "".join(f'<div class="card {k}"><span>{esc(k.upper())}</span><strong>{v:,}</strong></div>' for k, v in report["counts"].items() if k in {"processed", *STATUSES})
    body = []
    for row in report["rows"]:
        body.append(f'<tr data-status="{row["status"]}"><td>{row["source_row"]}</td><td><span class="badge {row["status"]}">{row["status"]}</span></td><td>{esc("; ".join(row["reasons"]) or "Passed all checks")}</td><td><details><summary>View original fields</summary><pre>{esc(json.dumps(row["original"], ensure_ascii=False, indent=2, default=serial))}</pre></details></td></tr>')
    document = '''<!doctype html><html lang="en" translate="no"><meta charset="utf-8"><meta name="google" content="notranslate"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Data quality audit</title><style>
*{box-sizing:border-box}body{margin:0;background:#f2f5f8;color:#173348;font:15px/1.5 system-ui,Segoe UI,sans-serif}main{max-width:1300px;margin:40px auto;padding:0 28px}header{border-bottom:1px solid #d3dfe7;padding-bottom:25px}.eyebrow{letter-spacing:.16em;font-size:12px;font-weight:750;color:#117c86}h1{font-size:38px;line-height:1.15;margin:10px 0}h2{font-size:21px;margin-top:34px}.muted{color:#617789;overflow-wrap:anywhere}.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(160px,1fr));gap:14px;margin:26px 0}.card{background:white;border:1px solid #d9e3ea;border-radius:12px;padding:19px;box-shadow:0 5px 18px #17334806}.card span{display:block;font-size:11px;letter-spacing:.1em;color:#617789}.card strong{display:block;font-size:34px;font-weight:650;margin-top:6px}.valid strong{color:#087b62}.invalid strong{color:#b74348}.duplicate strong{color:#956223}.filtered strong{color:#6250a4}.note{background:#e4f2f1;border-left:4px solid #198d89;padding:15px 20px;border-radius:5px}.controls{display:flex;justify-content:space-between;align-items:center;gap:16px;flex-wrap:wrap}select{background:white;border:1px solid #b9cbd8;border-radius:7px;padding:9px;color:#173348}.table{overflow:auto;border:1px solid #d9e3ea;border-radius:10px;background:white}table{width:100%;border-collapse:collapse}th{background:#173348;color:white;text-align:left;font-size:12px;letter-spacing:.04em;padding:14px}td{padding:13px 14px;border-bottom:1px solid #e4ebef;vertical-align:top}td:first-child{width:90px}td:nth-child(2){width:115px}td:nth-child(3){min-width:240px}pre{margin:0;font:12px/1.6 ui-monospace,Consolas,monospace;white-space:pre-wrap;overflow-wrap:anywhere;min-width:260px}.badge{display:inline-block;background:#e9eef3;border-radius:12px;padding:3px 10px;font-size:12px}.badge.valid{background:#dff3e9;color:#087b62}.badge.invalid{background:#fde8e8;color:#a43940}.badge.duplicate{background:#fff0d6;color:#956223}.badge.filtered{background:#eee8fb;color:#6250a4}footer{padding:25px 0;font-size:12px;color:#617789}@media(max-width:600px){main{margin:24px auto;padding:0 14px}h1{font-size:29px}}</style>
<main><header><div class="eyebrow">LOCAL DATA CLEANER / QUALITY CONTROL</div><h1>A clear record of every row.</h1><div class="muted">SOURCE_NAME<br>Generated GENERATED_AT</div></header><section class="cards">CARDS</section><div class="note"><b>Every source row is accounted for.</b> Processed = valid + invalid + duplicate + filtered. Blank rows skipped: BLANK_COUNT. Only valid rows appear in the clean CSV and workbook. Duplicate detection uses the first valid normalized key before filters.</div><div class="controls"><h2>Row audit</h2><label>Show status <select id="status"><option value="all">All rows</option><option>valid</option><option>invalid</option><option>duplicate</option><option>filtered</option></select></label></div><div class="table"><table><thead><tr><th>SOURCE ROW</th><th>STATUS</th><th>REASON</th><th>ORIGINAL VALUES</th></tr></thead><tbody>ROWS</tbody></table></div><footer>Processed locally. Source files are unchanged. Spreadsheet formulas are not evaluated. Report contains the original source values for review.</footer></main><script>document.getElementById('status').addEventListener('change',function(){document.querySelectorAll('tbody tr').forEach(r=>r.hidden=this.value!=='all'&&r.dataset.status!==this.value)});</script></html>'''
    document = document.replace("A clear record of every row.", "Excel and CSV cleaning audit")
    if Path(report["source"]).name == "9077_示例订单.csv":
        document = document.replace("</header>", '<p class="muted" style="font-size:12px;margin-bottom:0">Independent portfolio demo. Included sample data is synthetic.</p></header>', 1)
    document = document.replace("pre{margin:0;", "summary{cursor:pointer;color:#117c86;font-size:13px;white-space:nowrap}details[open] summary{margin-bottom:10px}pre{margin:0;")
    replacements = {"SOURCE_NAME": esc(Path(report["source"]).name), "GENERATED_AT": esc(report["generated_at"]), "CARDS": cards, "BLANK_COUNT": str(report["counts"]["blank_rows_skipped"]), "ROWS": "".join(body)}
    document = re.sub(r"\b(SOURCE_NAME|GENERATED_AT|CARDS|BLANK_COUNT|ROWS)\b", lambda match: replacements[match[0]], document)
    path.write_text(document, encoding="utf-8")


def run(input_path, config_path, output_dir):
    """Validate, clean, classify, and atomically publish into a new/empty folder."""
    config = load_config(config_path)
    raw_output = Path(output_dir).absolute()
    if raw_output.is_symlink():
        raise OutputError("Output directory must not be a symbolic link")
    destination = raw_output.resolve()
    if destination.exists() and (not destination.is_dir() or any(destination.iterdir())):
        raise OutputError("Output directory must be new or empty; existing files are never overwritten")
    _, rows, blank = read_source(input_path, config)
    valid = classify(rows, config)
    counts = Counter(row["status"] for row in rows)
    counts = {"processed": len(rows), **{key: counts[key] for key in STATUSES}, "blank_rows_skipped": blank}
    assert counts["processed"] == sum(counts[key] for key in STATUSES)
    report = {"schema_version": 1, "source": Path(input_path).name, "config": Path(config_path).name,
              "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "counts": counts,
              "outputs": dict(NAMES), "rows": json_value(rows)}
    stage = None
    try:
        destination.parent.mkdir(parents=True, exist_ok=True)
        stage = Path(tempfile.mkdtemp(prefix=PREFIX + "暂存_", dir=destination.parent))
        with (stage / NAMES["csv"]).open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow([safe_cell(name) for name in config["columns"]])
            for row in valid:
                writer.writerow([format(row["cleaned"][name], "f") if isinstance(row["cleaned"][name], Decimal) else safe_cell(row["cleaned"][name]) for name in config["columns"]])
        write_workbook(stage / NAMES["xlsx"], config["columns"], valid, report)
        (stage / NAMES["json"]).write_text(json.dumps(report, ensure_ascii=False, indent=2, default=serial, allow_nan=False), encoding="utf-8")
        write_html(stage / NAMES["html"], report)
        if destination.exists():
            destination.rmdir()  # Succeeds only if still empty; never deletes user files.
        os.replace(stage, destination)
        stage = None
        return json.loads(json.dumps(report, default=serial, allow_nan=False))
    except ImportError as exc:
        raise CleanerError("Missing dependency: install openpyxl with python -m pip install openpyxl") from exc
    except (OSError, ValueError) as exc:
        raise OutputError(f"Cannot publish outputs: {exc}") from exc
    finally:
        if stage is not None and stage.parent.resolve() == destination.parent and stage.name.startswith(PREFIX + "暂存_"):
            shutil.rmtree(stage)


def main(argv=None):
    parser = argparse.ArgumentParser(description="Clean and audit a local CSV/XLSX without modifying the source.")
    parser.add_argument("input", help="UTF-8 CSV or XLSX file")
    parser.add_argument("config", help="JSON configuration")
    parser.add_argument("output", help="new or empty output directory")
    args = parser.parse_args(argv)
    try:
        report = run(args.input, args.config, args.output)
        print(json.dumps({"counts": report["counts"], "outputs": report["outputs"]}, ensure_ascii=False, indent=2))
        return 0
    except CleanerError as exc:
        print(f"ERROR [{exc.exit_code}]: {exc}", file=sys.stderr)
        return exc.exit_code
    except Exception as exc:
        print(f"ERROR [5]: Unexpected failure: {exc}", file=sys.stderr)
        return 5


if __name__ == "__main__":
    raise SystemExit(main())
